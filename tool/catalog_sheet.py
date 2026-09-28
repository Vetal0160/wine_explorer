"""Каталог в таблице Excel/CSV ⇄ Supabase.

    python tool/catalog_sheet.py export [catalog.xlsx]
        Текущие винодельни и вина из Supabase → таблица Excel (листы
        «wineries», «wines» и «Инструкция»). Нужен supabase.json.

    python tool/catalog_sheet.py template [catalog.xlsx]
        Пустой шаблон той же структуры.

    python tool/catalog_sheet.py import <catalog.xlsx | папка с wineries.csv и wines.csv> [-o import.sql]
        Проверяет таблицу и готовит SQL: новые строки добавляются,
        существующие (винодельня — по названию, вино — по винодельне +
        названию + году) обновляются. Пустые ячейки не стирают то, что уже
        есть в базе. SQL вставляется в Supabase → SQL Editor → Run.

Колонки «sources» и «notes» в базу не попадают — это для ссылок и пометок.
"""
import csv
import datetime
import io
import json
import os
import re
import sys
import urllib.request

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LANGS = ("ru", "ro", "en")
WINE_TYPES = ("red_dry", "white_dry", "rose_dry", "sparkling")

# Типы, как их может написать человек или агент → код
TYPE_ALIASES = {
    "красное сухое": "red_dry", "белое сухое": "white_dry", "розовое сухое": "rose_dry",
    "розовое": "rose_dry", "игристое": "sparkling", "шампанское": "sparkling",
    "roșu sec": "red_dry", "rosu sec": "red_dry", "alb sec": "white_dry",
    "roze sec": "rose_dry", "rose sec": "rose_dry", "rosé sec": "rose_dry", "spumant": "sparkling",
    "red dry": "red_dry", "dry red": "red_dry", "red": "red_dry",
    "white dry": "white_dry", "dry white": "white_dry", "white": "white_dry",
    "rose dry": "rose_dry", "rosé dry": "rose_dry", "dry rosé": "rose_dry", "rosé": "rose_dry",
    "rose": "rose_dry", "sparkling": "sparkling",
}
TRUE_WORDS = {"да", "yes", "y", "true", "1", "da", "+"}
FALSE_WORDS = {"нет", "no", "n", "false", "0", "nu", "-"}

# (ключ, обязательная?, подсказка) — порядок колонок в таблице
WINERY_COLS = [
    ("name", True, "Название, как на этикетке: «Château Purcari». По нему обновляется существующая винодельня."),
    ("region", False, "Регион · село: «Ștefan Vodă · Purcari»"),
    ("latitude", True, "Широта, число с точкой: 46.5250 (Google Карты → правый клик по месту)"),
    ("longitude", True, "Долгота: 29.8650"),
    ("founded_year", False, "Год основания: 1827"),
    ("address", False, "Адрес для посетителей"),
    ("phone", False, "Телефон в формате +373 …"),
    ("website", False, "Сайт: https://…"),
    ("hours_ru", False, "Часы работы по-русски: «Пн–Сб 10:00–18:00»"),
    ("hours_ro", False, "Часы по-румынски: «Lun–Sâm 10:00–18:00»"),
    ("hours_en", False, "Часы по-английски: «Mon–Sat 10:00–18:00»"),
    ("has_tastings", False, "Есть дегустации: да / нет (пусто — неизвестно)"),
    ("tasting_price_lei", False, "Цена дегустации от, MDL: 250"),
    ("booking_required", False, "Нужна запись: да / нет"),
    ("description_ru", False, "Описание, 1–3 предложения, по-русски"),
    ("description_ro", False, "Описание по-румынски"),
    ("description_en", False, "Описание по-английски"),
    ("image_url", False, "Фото: путь в Storage («purcari/winery.jpg») или https://…"),
    ("sources", False, "Откуда данные (ссылки) — в базу не попадает"),
    ("notes", False, "Пометки для себя — в базу не попадает"),
]
WINE_COLS = [
    ("name", True, "Название вина: «Viorica de Purcari»"),
    ("winery", True, "Винодельня — точно как в листе wineries (или уже в базе)"),
    ("type", True, "Тип: red_dry / white_dry / rose_dry / sparkling (можно «Красное сухое»)"),
    ("grape_variety", False, "Сорт(а) через запятую: «Chardonnay, Pinot Noir»"),
    ("vintage", False, "Год урожая: 2021 (пусто или NV — без года)"),
    ("rating", False, "Оценка 0–5: 4.6 (пусто — без оценки)"),
    ("price_lei", False, "Примерная цена в магазине, MDL: 180"),
    ("description_ru", False, "Описание вкуса по-русски"),
    ("description_ro", False, "Описание по-румынски"),
    ("description_en", False, "Описание по-английски"),
    ("image_url", False, "Фото: путь в Storage («purcari/viorica.jpg») или https://…"),
    ("sources", False, "Откуда данные (ссылки) — в базу не попадает"),
    ("notes", False, "Пометки для себя — в базу не попадает"),
]


# ---------------------------------------------------------------- Supabase (только чтение)

def supabase_config():
    path = os.path.join(ROOT, "supabase.json")
    if not os.path.exists(path):
        return None
    cfg = json.load(open(path, encoding="utf-8"))
    return cfg["SUPABASE_URL"].rstrip("/"), cfg["SUPABASE_KEY"]


def rest_get(query):
    cfg = supabase_config()
    if cfg is None:
        raise SystemExit("Нет supabase.json — не могу прочитать данные из Supabase.")
    url, key = cfg
    req = urllib.request.Request(f"{url}/rest/v1/{query}", headers={"apikey": key})
    with urllib.request.urlopen(req, timeout=30) as resp:
        return json.loads(resp.read().decode("utf-8"))


# ---------------------------------------------------------------- Excel

HEADER_FILL = PatternFill("solid", fgColor="58111A")
REQUIRED_FILL = PatternFill("solid", fgColor="8B263E")


def write_sheet(ws, cols, rows):
    for i, (key, required, hint) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=key)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = REQUIRED_FILL if required else HEADER_FILL
        cell.comment = Comment(("ОБЯЗАТЕЛЬНО. " if required else "") + hint, "Wine Explorer")
        wide = key.startswith("description") or key in ("sources", "notes", "address")
        ws.column_dimensions[cell.column_letter].width = 45 if wide else 20
    for r, row in enumerate(rows, start=2):
        for i, (key, _, _) in enumerate(cols, start=1):
            value = row.get(key)
            ws.cell(row=r, column=i, value=value).alignment = Alignment(wrap_text=key.startswith("description"), vertical="top")
    ws.freeze_panes = "B2"


def add_list_validation(ws, cols, key, values, max_row=2000):
    idx = [c[0] for c in cols].index(key) + 1
    letter = ws.cell(row=1, column=idx).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error = "Выберите значение из списка"
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{max_row}")


def instruction_sheet(ws):
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    lines = [
        ("Как пользоваться", ""),
        ("1.", "Заполняйте листы «wineries» (винодельни) и «wines» (вина). Одна строка — одна запись."),
        ("2.", "Тёмно-красные заголовки — обязательные колонки. Наведите на заголовок — будет подсказка."),
        ("3.", "Сначала винодельня, потом её вина: колонка wines.winery должна совпадать с wineries.name."),
        ("4.", "Существующие записи обновляются: винодельня — по названию, вино — по винодельне + названию + году."),
        ("5.", "Пустая ячейка не стирает данные в базе. Чтобы удалить значение — правьте в Supabase."),
        ("6.", "Импорт: python tool/catalog_sheet.py import catalog.xlsx → вставить SQL в Supabase → SQL Editor → Run."),
        ("", ""),
        ("Правила", ""),
        ("Факты", "Только проверенные: телефон, часы, цены — с сайта винодельни или магазина. Не знаете — оставьте пусто."),
        ("Тексты", "Описания — своими словами, не копировать с сайтов (авторское право)."),
        ("Фото", "Только свои или с разрешения винодельни. Чужие фото с сайтов не вставлять."),
        ("Числа", "Точка или запятая — обе подходят: 4.6 и 4,6."),
        ("", ""),
        ("Колонки: винодельни", ""),
        *[(k, ("ОБЯЗАТЕЛЬНО. " if req else "") + h) for k, req, h in WINERY_COLS],
        ("", ""),
        ("Колонки: вина", ""),
        *[(k, ("ОБЯЗАТЕЛЬНО. " if req else "") + h) for k, req, h in WINE_COLS],
    ]
    for r, (a, b) in enumerate(lines, start=1):
        ws.cell(row=r, column=1, value=a).font = Font(bold=not b or a in ("Факты", "Тексты", "Фото", "Числа"))
        ws.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True)


def save_workbook(path, winery_rows, wine_rows):
    wb = Workbook()
    ws_w = wb.active
    ws_w.title = "wineries"
    write_sheet(ws_w, WINERY_COLS, winery_rows)
    add_list_validation(ws_w, WINERY_COLS, "has_tastings", ["да", "нет"])
    add_list_validation(ws_w, WINERY_COLS, "booking_required", ["да", "нет"])
    ws_v = wb.create_sheet("wines")
    write_sheet(ws_v, WINE_COLS, wine_rows)
    add_list_validation(ws_v, WINE_COLS, "type", list(WINE_TYPES))
    instruction_sheet(wb.create_sheet("Инструкция"))
    wb.save(path)


def export(path):
    wineries = rest_get("wineries?select=*&order=name.asc")
    wines = rest_get("wines?select=*,winery:wineries(name)&order=winery_id.asc,name.asc")
    bool_word = {True: "да", False: "нет", None: None}
    winery_rows = []
    for w in wineries:
        row = {k: w.get(k) for k in ("name", "region", "latitude", "longitude", "founded_year",
                                     "address", "phone", "website", "tasting_price_lei", "image_url")}
        row["has_tastings"] = bool_word[w.get("has_tastings")]
        row["booking_required"] = bool_word[w.get("booking_required")]
        for lang in LANGS:
            row[f"description_{lang}"] = (w.get("description") or {}).get(lang)
            row[f"hours_{lang}"] = (w.get("hours") or {}).get(lang)
        winery_rows.append(row)
    wine_rows = []
    for v in wines:
        row = {"name": v["name"], "winery": (v.get("winery") or {}).get("name"), "type": v["type"],
               "grape_variety": v.get("grape_variety"), "vintage": v.get("vintage"),
               "rating": v.get("rating"), "price_lei": v.get("avg_price_lei"), "image_url": v.get("image_url")}
        for lang in LANGS:
            row[f"description_{lang}"] = (v.get("description") or {}).get(lang)
        wine_rows.append(row)
    save_workbook(path, winery_rows, wine_rows)
    print(f"Готово: {path} — виноделен: {len(winery_rows)}, вин: {len(wine_rows)}")


# ---------------------------------------------------------------- Чтение таблицы

def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_xlsx(path):
    wb = load_workbook(path, data_only=True)
    sheets = {}
    for name in ("wineries", "wines"):
        if name not in wb.sheetnames:
            raise SystemExit(f"В файле нет листа «{name}».")
        rows = list(wb[name].iter_rows(values_only=True))
        header = [clean(h) for h in rows[0]] if rows else []
        sheets[name] = [
            (i, {h: clean(v) for h, v in zip(header, r) if h})
            for i, r in enumerate(rows[1:], start=2)
            if any(clean(v) for v in r)
        ]
    return sheets["wineries"], sheets["wines"]


def read_csv(path):
    raw = open(path, "rb").read().decode("utf-8-sig")
    dialect = csv.Sniffer().sniff(raw.split("\n", 1)[0], delimiters=",;\t")
    reader = csv.DictReader(io.StringIO(raw), dialect=dialect)
    return [
        (i, {clean(k): clean(v) for k, v in row.items() if k})
        for i, row in enumerate(reader, start=2)
        if any(clean(v) for v in row.values())
    ]


def read_input(path):
    full = os.path.abspath(path)
    if not os.path.exists(full):
        raise SystemExit(
            f"Не найдено: {full}\n"
            "Создайте эту папку и положите в неё wineries.csv и wines.csv "
            "(или укажите путь к файлу .xlsx).")
    if os.path.isdir(full):
        files = {f.lower(): f for f in os.listdir(full)}
        found = {name: os.path.join(full, files[name]) for name in ("wineries.csv", "wines.csv") if name in files}
        if not found:
            listing = ", ".join(sorted(os.listdir(full))) or "папка пустая"
            raise SystemExit(
                f"В папке {full} нет файлов wineries.csv и wines.csv.\n"
                f"Сейчас в ней: {listing}\n"
                "Переименуйте файлы агента точно так: wineries.csv и wines.csv.")
        for name in ("wineries.csv", "wines.csv"):
            if name not in found:
                print(f"  ! нет файла {name} — загружаю только то, что есть")
        return (read_csv(found["wineries.csv"]) if "wineries.csv" in found else [],
                read_csv(found["wines.csv"]) if "wines.csv" in found else [])
    if full.lower().endswith(".xlsx"):
        return read_xlsx(full)
    raise SystemExit(f"Не понимаю этот файл: {full}\nНужен .xlsx или папка с wineries.csv и wines.csv.")


# ---------------------------------------------------------------- Проверка

class Report:
    def __init__(self):
        self.errors, self.warnings = [], []

    def error(self, where, msg):
        self.errors.append(f"  ✗ {where}: {msg}")

    def warn(self, where, msg):
        self.warnings.append(f"  ! {where}: {msg}")


def parse_number(text):
    t = text.replace(" ", " ").strip()
    t = re.sub(r"(?i)\s*(mdl|lei|лей|леев|леи)\.?$", "", t).replace(" ", "").replace(",", ".")
    return float(t)


def parse_bool(text):
    t = text.strip().lower()
    if t in TRUE_WORDS:
        return True
    if t in FALSE_WORDS:
        return False
    raise ValueError(text)


def localized(row, prefix):
    return {lang: row[f"{prefix}_{lang}"] for lang in LANGS if row.get(f"{prefix}_{lang}")}


def validate_wineries(rows, rep):
    out, seen = [], set()
    for line, r in rows:
        where = f"wineries, строка {line}"
        name = r.get("name", "")
        if not name:
            rep.error(where, "пустое название (name)")
            continue
        if name.lower() in seen:
            rep.error(where, f"«{name}» уже есть выше — одна винодельня должна быть в одной строке")
            continue
        seen.add(name.lower())
        w = {"name": name, "region": r.get("region") or None, "address": r.get("address") or None,
             "phone": r.get("phone") or None, "website": r.get("website") or None,
             "image_url": r.get("image_url", ""),
             "description": localized(r, "description"), "hours": localized(r, "hours")}
        try:
            w["latitude"] = parse_number(r.get("latitude", ""))
            w["longitude"] = parse_number(r.get("longitude", ""))
            if not (45.4 <= w["latitude"] <= 48.6 and 26.6 <= w["longitude"] <= 30.2):
                rep.warn(where, f"координаты {w['latitude']}, {w['longitude']} вне Молдовы — проверьте (не перепутаны ли широта и долгота?)")
        except ValueError:
            rep.error(where, f"«{name}»: широта и долгота обязательны и должны быть числами (46.5250)")
            continue
        for key, lo, hi in (("founded_year", 1000, 2100),):
            if r.get(key):
                try:
                    w[key] = int(parse_number(r[key]))
                    if not lo <= w[key] <= hi:
                        raise ValueError
                except ValueError:
                    rep.error(where, f"«{name}»: {key} — год числом, например 1827")
            else:
                w[key] = None
        if r.get("tasting_price_lei"):
            try:
                w["tasting_price_lei"] = parse_number(r["tasting_price_lei"])
            except ValueError:
                rep.error(where, f"«{name}»: tasting_price_lei — число, например 250")
        else:
            w["tasting_price_lei"] = None
        for key in ("has_tastings", "booking_required"):
            if r.get(key):
                try:
                    w[key] = parse_bool(r[key])
                except ValueError:
                    rep.error(where, f"«{name}»: {key} — «да» или «нет», а не «{r[key]}»")
            else:
                w[key] = None
        if w["website"] and not re.match(r"https?://", w["website"]):
            w["website"] = "https://" + w["website"]
        if w["phone"] and not w["phone"].startswith("+"):
            rep.warn(where, f"«{name}»: телефон лучше в формате +373 …")
        if not w["description"]:
            rep.warn(where, f"«{name}»: нет описания")
        elif len(w["description"]) < 3:
            rep.warn(where, f"«{name}»: описание не на всех языках ({', '.join(w['description'])})")
        out.append(w)
    return out


def validate_wines(rows, known_wineries, grape_names, rep):
    out, seen = [], set()
    this_year = datetime.date.today().year
    for line, r in rows:
        where = f"wines, строка {line}"
        name, winery = r.get("name", ""), r.get("winery", "")
        if not name or not winery:
            rep.error(where, "обязательны name и winery")
            continue
        match = known_wineries.get(winery.lower())
        if match is None:
            rep.error(where, f"«{name}»: винодельни «{winery}» нет ни в листе wineries, ни в базе")
            continue
        raw_type = r.get("type", "").strip()
        wine_type = raw_type if raw_type in WINE_TYPES else TYPE_ALIASES.get(raw_type.lower())
        if wine_type is None:
            rep.error(where, f"«{name}»: тип «{raw_type}» — нужен один из {', '.join(WINE_TYPES)}")
            continue
        v = {"name": name, "winery": match, "type": wine_type,
             "grape_variety": r.get("grape_variety", ""), "image_url": r.get("image_url", ""),
             "description": localized(r, "description")}
        vintage = r.get("vintage", "").strip()
        if vintage and vintage.upper() not in ("NV", "N/V", "-"):
            try:
                v["vintage"] = int(parse_number(vintage))
                if not 1900 <= v["vintage"] <= this_year + 1:
                    raise ValueError
            except ValueError:
                rep.error(where, f"«{name}»: год урожая «{vintage}» — число вроде 2021 или NV")
                continue
        else:
            v["vintage"] = None
        for key, target, lo, hi in (("rating", "rating", 0, 5), ("price_lei", "price", 0, 100000)):
            if r.get(key):
                try:
                    v[target] = parse_number(r[key])
                    if not lo <= v[target] <= hi:
                        raise ValueError
                except ValueError:
                    rep.error(where, f"«{name}»: {key} «{r[key]}» — число от {lo} до {hi}")
                    v[target] = None
            else:
                v[target] = None
        key = (match.lower(), name.lower(), v["vintage"])
        if key in seen:
            rep.error(where, f"«{name}» {v['vintage'] or 'NV'} от «{match}» уже есть выше")
            continue
        seen.add(key)
        if grape_names:
            for g in re.split(r"[,/&+;]", v["grape_variety"]):
                if g.strip() and normalize(g) not in grape_names:
                    rep.warn(where, f"«{name}»: сорта «{g.strip()}» нет в справочнике — связь со страницей сорта не появится")
        if not v["description"]:
            rep.warn(where, f"«{name}»: нет описания")
        out.append(v)
    return out


def normalize(name):
    table = str.maketrans("ăâîșşțţéèô", "aaisstteeo")
    return re.sub(r"\s+", " ", name.strip().lower().translate(table))


# ---------------------------------------------------------------- SQL

def q(value):
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def q_json(obj):
    return q(json.dumps(obj, ensure_ascii=False)) + "::jsonb"


def q_num(value, cast):
    return f"null::{cast}" if value is None else f"{value}::{cast}"


def q_bool(value):
    return "null" if value is None else ("true" if value else "false")


def build_sql(wineries, wines, source_name):
    out = [
        f"-- Импорт каталога из «{source_name}», {datetime.datetime.now():%Y-%m-%d %H:%M}.",
        f"-- Виноделен: {len(wineries)}, вин: {len(wines)}. Supabase → SQL Editor → Run.",
        "-- Всё в одной транзакции: при ошибке ничего не изменится.",
        "begin;",
        "",
        "-- Вино определяется винодельней + названием + годом (нужно для обновления)",
        "do $$",
        "begin",
        "  if not exists (select 1 from pg_constraint where conname = 'wines_winery_name_vintage_key') then",
        "    alter table public.wines",
        "      add constraint wines_winery_name_vintage_key unique nulls not distinct (winery_id, name, vintage);",
        "  end if;",
        "end $$;",
        "",
    ]
    if wineries:
        out.append("insert into public.wineries (name, region, latitude, longitude, founded_year, address, phone, website,")
        out.append("  hours, has_tastings, tasting_price_lei, booking_required, description, image_url) values")
        rows = []
        for w in wineries:
            rows.append(
                f"  ({q(w['name'])}, {q(w['region'])}, {w['latitude']}, {w['longitude']}, "
                f"{q_num(w['founded_year'], 'integer')}, {q(w['address'])}, {q(w['phone'])}, {q(w['website'])},\n"
                f"   {q_json(w['hours'])}, {q_bool(w['has_tastings'])}, {q_num(w['tasting_price_lei'], 'numeric')}, "
                f"{q_bool(w['booking_required'])},\n   {q_json(w['description'])}, {q(w['image_url'])})")
        out.append(",\n".join(rows))
        out += [
            "on conflict (name) do update set",
            "  region            = coalesce(excluded.region, wineries.region),",
            "  latitude          = excluded.latitude,",
            "  longitude         = excluded.longitude,",
            "  founded_year      = coalesce(excluded.founded_year, wineries.founded_year),",
            "  address           = coalesce(excluded.address, wineries.address),",
            "  phone             = coalesce(excluded.phone, wineries.phone),",
            "  website           = coalesce(excluded.website, wineries.website),",
            "  hours             = wineries.hours || excluded.hours,",
            "  has_tastings      = coalesce(excluded.has_tastings, wineries.has_tastings),",
            "  tasting_price_lei = coalesce(excluded.tasting_price_lei, wineries.tasting_price_lei),",
            "  booking_required  = coalesce(excluded.booking_required, wineries.booking_required),",
            "  description       = wineries.description || excluded.description,",
            "  image_url         = case when excluded.image_url <> '' then excluded.image_url else wineries.image_url end;",
            "",
        ]
    if wines:
        out.append("insert into public.wines (winery_id, name, type, grape_variety, vintage, rating, avg_price_lei, image_url, description)")
        out.append("select w.id, v.name, v.type, v.grape_variety, v.vintage, v.rating, v.price, v.image_url, v.description")
        out.append("from (values")
        rows = []
        for v in wines:
            rows.append(
                f"  ({q(v['winery'])}::text, {q(v['name'])}::text, {q(v['type'])}::text, {q(v['grape_variety'])}::text, "
                f"{q_num(v['vintage'], 'integer')}, {q_num(v['rating'], 'numeric')}, {q_num(v['price'], 'numeric')}, "
                f"{q(v['image_url'])}::text,\n   {q_json(v['description'])})")
        out.append(",\n".join(rows))
        out += [
            ") as v (winery, name, type, grape_variety, vintage, rating, price, image_url, description)",
            "join public.wineries w on w.name = v.winery",
            "on conflict on constraint wines_winery_name_vintage_key do update set",
            "  type          = excluded.type,",
            "  grape_variety = case when excluded.grape_variety <> '' then excluded.grape_variety else wines.grape_variety end,",
            "  rating        = coalesce(excluded.rating, wines.rating),",
            "  avg_price_lei = coalesce(excluded.avg_price_lei, wines.avg_price_lei),",
            "  image_url     = case when excluded.image_url <> '' then excluded.image_url else wines.image_url end,",
            "  description   = wines.description || excluded.description;",
            "",
        ]
    out += ["commit;", ""]
    return "\n".join(out)


def do_import(path, out_path):
    winery_rows, wine_rows = read_input(path)
    rep = Report()
    wineries = validate_wineries(winery_rows, rep)

    # Винодельни и сорта, уже лежащие в базе (для проверки ссылок); без сети — только из таблицы
    known = {w["name"].lower(): w["name"] for w in wineries}
    grape_names = set()
    if supabase_config():
        try:
            for w in rest_get("wineries?select=name"):
                known.setdefault(w["name"].lower(), w["name"])
            for g in rest_get("grapes?select=name,aliases"):
                grape_names.add(normalize(g["name"]))
                grape_names.update(normalize(a) for a in g.get("aliases") or [])
        except Exception as e:  # noqa: BLE001 — проверка ссылок не критична
            rep.warn("Supabase", f"не удалось прочитать базу ({e}); проверяю только по таблице")
    wines = validate_wines(wine_rows, known, grape_names, rep)

    for line in rep.warnings:
        print(line)
    if rep.errors:
        print(f"\nОшибок: {len(rep.errors)} — исправьте и запустите снова. SQL не создан.")
        for line in rep.errors:
            print(line)
        sys.exit(1)
    if not wineries and not wines:
        raise SystemExit("В таблице нет строк для импорта.")

    sql = build_sql(wineries, wines, os.path.basename(path.rstrip("/\\")))
    open(out_path, "w", encoding="utf-8", newline="\n").write(sql)
    print(f"\nГотово: {out_path}\nВиноделен: {len(wineries)}, вин: {len(wines)}, предупреждений: {len(rep.warnings)}")
    print("Дальше: откройте файл, скопируйте всё → Supabase → SQL Editor → New query → Run.")


def main():
    for stream in (sys.stdout, sys.stderr):  # кириллица в консоли Windows
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    args = sys.argv[1:]
    if not args or args[0] not in ("export", "template", "import"):
        print(__doc__)
        sys.exit(1)
    cmd = args[0]
    if cmd == "import":
        if len(args) < 2:
            raise SystemExit("Укажите таблицу: python tool/catalog_sheet.py import catalog.xlsx")
        out = args[args.index("-o") + 1] if "-o" in args else os.path.join(
            os.path.dirname(os.path.abspath(args[1])), "import.sql")
        do_import(args[1], out)
    else:
        path = args[1] if len(args) > 1 else "catalog.xlsx"
        if cmd == "export":
            export(path)
        else:
            save_workbook(path, [], [])
            print(f"Готово: {path}")


if __name__ == "__main__":
    main()
